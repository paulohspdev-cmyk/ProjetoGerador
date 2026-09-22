import hashlib
import json
import os
import sqlite3
import tempfile
import threading
import time
from pathlib import Path

from openpyxl import load_workbook

root_tmp = tempfile.TemporaryDirectory(prefix="rc-audit-reg-")
root = Path(root_tmp.name)
data = root / "data"
archive = root / "rapid-archive"
archive.mkdir(parents=True)
(archive / "history.bin").write_bytes(b"history")
os.environ["RC_DATA_DIR"] = str(data)
os.environ["RC_DB_FILE"] = str(data / "rc-geradores.db")
scada_root = root / "scada"
for rel in ("BaseDAT", "Config", "ScadaComm/Config"):
    target = scada_root / rel
    target.mkdir(parents=True, exist_ok=True)
    (target / "placeholder.txt").write_text("test")
os.environ["RC_RAPID_SCADA_ROOT"] = str(scada_root)
os.environ["RC_TOTP_KEY_FILE"] = str(root / "totp.key")
os.environ["RC_RAPID_ARCHIVE_DIR"] = str(archive)
os.environ["RC_PUBLIC_BASE_URL"] = "https://example.invalid"
os.environ["RC_SMTP_HOST"] = "smtp.invalid"
os.environ["RC_SMTP_FROM"] = "noreply@example.invalid"

from app import (  # noqa: E402
    automation_engine,
    db,
    diagnostics,
    industrial_store,
    ops_store,
    platform_store,
    traffic_store,
)
from app.auth import hash_password  # noqa: E402
from app.completion_routes import (  # noqa: E402
    _finish_staged_file,
    _restore_staged_file,
    _stage_file_delete,
)
from app.extra_routes import _token_allows_generator  # noqa: E402
from app.rapid import _downsample_points, dashboard  # noqa: E402
from app.migrations import _generator_nominal_power_v3, _operator_role_v2  # noqa: E402
from app.reporting import generate_report, safe_report_artifact_path  # noqa: E402
from app.secret_box import protect_secret  # noqa: E402
from app.security_service import disable_totp, setup_totp, totp_code  # noqa: E402


def init_all():
    db.init_db()
    ops_store.init_ops_db()
    platform_store.init_platform_db()
    industrial_store.init_industrial_db()


init_all()

# F00: external TLS proxy mode must not report local nginx as a failed product service,
# but readiness must surface a legacy local TLS terminator that remains active on 443.
previous_web_tls_mode = os.environ.get("RC_WEB_TLS_MODE")
previous_trusted_proxy_cidrs = os.environ.get("RC_TRUSTED_PROXY_CIDRS")
_original_proxy_run = diagnostics._run
try:
    os.environ["RC_WEB_TLS_MODE"] = "external_proxy"
    os.environ["RC_TRUSTED_PROXY_CIDRS"] = ""
    assert "nginx.service" not in diagnostics._service_names()

    def _local_nginx_active(args, timeout=2):
        if args == ["systemctl", "is-active", "nginx.service"]:
            return 0, "active"
        return _original_proxy_run(args, timeout)

    diagnostics._run = _local_nginx_active
    proxy_ok, proxy_detail = diagnostics._external_proxy_topology({80, 443})
    assert proxy_ok is False, proxy_detail
    assert "dupla terminação TLS" in proxy_detail, proxy_detail

    proxy_readiness = diagnostics._production_readiness(
        [],
        reverse_tcp_exposed=False,
        reverse_tcp_allowlist=False,
        external_proxy_topology_ok=proxy_ok,
        external_proxy_topology_detail=proxy_detail,
    )
    proxy_check = next(
        item
        for item in proxy_readiness["checks"]
        if item["id"] == "external_proxy_topology"
    )
    assert proxy_check["severity"] == "blocker", proxy_check
    trusted_proxy_check = next(
        item
        for item in proxy_readiness["checks"]
        if item["id"] == "trusted_proxy_identity"
    )
    assert trusted_proxy_check["severity"] == "blocker", trusted_proxy_check
    assert trusted_proxy_check["ok"] is False, trusted_proxy_check

    os.environ["RC_TRUSTED_PROXY_CIDRS"] = "10.10.10.131/32"
    trusted_readiness = diagnostics._production_readiness(
        [],
        reverse_tcp_exposed=False,
        reverse_tcp_allowlist=False,
        external_proxy_topology_ok=True,
        external_proxy_topology_detail="proxy externo sem TLS local",
    )
    trusted_proxy_check = next(
        item
        for item in trusted_readiness["checks"]
        if item["id"] == "trusted_proxy_identity"
    )
    assert trusted_proxy_check["ok"] is True, trusted_proxy_check

    proxy_ok, proxy_detail = diagnostics._external_proxy_topology({3000, 8090})
    assert proxy_ok is True, proxy_detail

    os.environ["RC_WEB_TLS_MODE"] = "managed"
    assert "nginx.service" in diagnostics._service_names()
    proxy_ok, proxy_detail = diagnostics._external_proxy_topology({443})
    assert proxy_ok is None and proxy_detail == ""
finally:
    diagnostics._run = _original_proxy_run
    if previous_web_tls_mode is None:
        os.environ.pop("RC_WEB_TLS_MODE", None)
    else:
        os.environ["RC_WEB_TLS_MODE"] = previous_web_tls_mode
    if previous_trusted_proxy_cidrs is None:
        os.environ.pop("RC_TRUSTED_PROXY_CIDRS", None)
    else:
        os.environ["RC_TRUSTED_PROXY_CIDRS"] = previous_trusted_proxy_cidrs

# F00b2: readiness must verify the systemd policy actually loaded for native Rapid ports.
_previous_admin_cidrs = os.environ.get("RC_RAPID_ADMIN_ALLOWED_CIDRS")
_original_diagnostics_run = diagnostics._run
try:
    os.environ["RC_RAPID_ADMIN_ALLOWED_CIDRS"] = "10.10.10.0/24"

    def _rapid_policy_ok_run(args, timeout=2):
        if args[:2] == ["systemctl", "show"]:
            return (
                0,
                "IPAddressAllow=10.10.10.0/24 127.0.0.0/8 ::1/128\n"
                "IPAddressDeny=0.0.0.0/0 ::/0",
            )
        return _original_diagnostics_run(args, timeout)

    diagnostics._run = _rapid_policy_ok_run
    policy_ok, policy_detail = diagnostics._rapid_native_network_policy()
    assert policy_ok is True, policy_detail
    assert "10.10.10.0/24" in policy_detail, policy_detail

    os.environ["RC_RAPID_ADMIN_ALLOWED_CIDRS"] = ""

    def _rapid_policy_loopback_only_run(args, timeout=2):
        if args[:2] == ["systemctl", "show"]:
            return (
                0,
                "IPAddressAllow=127.0.0.0/8 ::1/128\n"
                "IPAddressDeny=0.0.0.0/0 ::/0",
            )
        return _original_diagnostics_run(args, timeout)

    diagnostics._run = _rapid_policy_loopback_only_run
    policy_ok, policy_detail = diagnostics._rapid_native_network_policy()
    assert policy_ok is True, policy_detail
    assert "somente a loopback" in policy_detail, policy_detail

    os.environ["RC_RAPID_ADMIN_ALLOWED_CIDRS"] = "10.10.10.0/24"

    def _rapid_policy_missing_deny(args, timeout=2):
        if args[:2] == ["systemctl", "show"]:
            service = args[2]
            deny = "" if service == "scadaweb6.service" else "IPAddressDeny=0.0.0.0/0 ::/0"
            return (
                0,
                "IPAddressAllow=10.10.10.0/24 127.0.0.0/8 ::1/128\n" + deny,
            )
        return _original_diagnostics_run(args, timeout)

    diagnostics._run = _rapid_policy_missing_deny
    policy_ok, policy_detail = diagnostics._rapid_native_network_policy()
    assert policy_ok is False, policy_detail
    assert "scadaweb6.service" in policy_detail, policy_detail
    assert "deny ausente" in policy_detail, policy_detail

    os.environ["RC_RAPID_ADMIN_ALLOWED_CIDRS"] = "0.0.0.0/0"
    policy_ok, policy_detail = diagnostics._rapid_native_network_policy()
    assert policy_ok is False, policy_detail
    assert "ampla demais" in policy_detail, policy_detail

    policy_readiness = diagnostics._production_readiness(
        [],
        reverse_tcp_exposed=False,
        reverse_tcp_allowlist=False,
        rapid_native_policy_ok=False,
        rapid_native_policy_detail="scadaweb6.service: deny ausente 0.0.0.0/0",
    )
    policy_check = next(
        item
        for item in policy_readiness["checks"]
        if item["id"] == "rapid_native_network_policy"
    )
    assert policy_check["severity"] == "blocker", policy_check
    assert "deny ausente" in policy_check["detail"], policy_check
finally:
    diagnostics._run = _original_diagnostics_run
    if _previous_admin_cidrs is None:
        os.environ.pop("RC_RAPID_ADMIN_ALLOWED_CIDRS", None)
    else:
        os.environ["RC_RAPID_ADMIN_ALLOWED_CIDRS"] = _previous_admin_cidrs

# F00c: readiness must not demand a nominal-power field that the inventory cannot store.
readiness = diagnostics._production_readiness(
    [
        {
            "id": "readiness-ig200",
            "tag": "READY-IG200",
            "controller_model": "InteliGen 200",
            "enabled": True,
            "site": "Lab",
            "customer": "Cliente",
        }
    ],
    reverse_tcp_exposed=False,
    reverse_tcp_allowlist=False,
)
nominal_check = next(
    item for item in readiness["checks"] if item["id"] == "nominal_power"
)
assert nominal_check["severity"] == "ok", nominal_check
assert "telemetria" in nominal_check["detail"].lower(), nominal_check

# F00d: firmware desconhecido só bloqueia packs que podem emitir comando industrial.
original_load_bindings = diagnostics.load_bindings
original_list_assets = diagnostics.domain_store.list_assets
original_list_controllers = diagnostics.domain_store.list_controllers
try:
    diagnostics.load_bindings = lambda: [
        {"generator_id": "fw-command"},
        {"generator_id": "fw-readonly"},
    ]
    diagnostics.domain_store.list_assets = lambda: [
        {"id": "asset-command", "legacy_generator_id": "fw-command"},
        {"id": "asset-readonly", "legacy_generator_id": "fw-readonly"},
    ]
    diagnostics.domain_store.list_controllers = lambda: [
        {"id": "ctrl-command", "asset_id": "asset-command", "firmware": ""},
        {"id": "ctrl-readonly", "asset_id": "asset-readonly", "firmware": ""},
    ]
    firmware_readiness = diagnostics._production_readiness(
        [
            {
                "id": "fw-command",
                "tag": "FW-CMD",
                "controller_model": "InteliGen 200",
                "enabled": True,
                "site": "Usina",
                "customer": "Cliente",
            },
            {
                "id": "fw-readonly",
                "tag": "FW-READ",
                "controller_model": "DSE4520 MKII",
                "enabled": True,
                "site": "Usina",
                "customer": "Cliente",
                "nominal_power_kw": 450.0,
            },
            {
                "id": "fw-no-pack",
                "tag": "FW-NOPACK",
                "controller_model": "InteliCompact NT",
                "enabled": True,
                "site": "Usina",
                "customer": "Cliente",
            },
        ],
        reverse_tcp_exposed=False,
        reverse_tcp_allowlist=False,
    )
finally:
    diagnostics.load_bindings = original_load_bindings
    diagnostics.domain_store.list_assets = original_list_assets
    diagnostics.domain_store.list_controllers = original_list_controllers

command_firmware = next(
    item for item in firmware_readiness["checks"] if item["id"] == "controller_firmware"
)
readonly_firmware = next(
    item
    for item in firmware_readiness["checks"]
    if item["id"] == "controller_firmware_readonly"
)
assert command_firmware["severity"] == "blocker", command_firmware
assert "FW-CMD" in command_firmware["detail"], command_firmware
assert "FW-READ" not in command_firmware["detail"], command_firmware
assert readonly_firmware["severity"] == "warning", readonly_firmware
assert "FW-READ" in readonly_firmware["detail"], readonly_firmware
assert "FW-CMD" not in readonly_firmware["detail"], readonly_firmware
assert "FW-NOPACK" not in command_firmware["detail"], command_firmware
assert "FW-NOPACK" not in readonly_firmware["detail"], readonly_firmware
nominal_readiness = next(
    item for item in firmware_readiness["checks"] if item["id"] == "nominal_power"
)
assert "FW-NOPACK" not in nominal_readiness["detail"], nominal_readiness
assert "FW-READ" not in nominal_readiness["detail"], nominal_readiness

# F00e: migration v3 adds nullable cadastral nominal power without fabricating a value.
legacy_nominal_path = root / "legacy-nominal.sqlite3"
legacy_nominal = sqlite3.connect(legacy_nominal_path)
legacy_nominal.execute(
    """CREATE TABLE generators(
        id TEXT PRIMARY KEY,
        tag TEXT NOT NULL,
        enabled INTEGER NOT NULL DEFAULT 1
    )"""
)
legacy_nominal.execute("INSERT INTO generators(id,tag,enabled) VALUES ('g1','LEGACY',1)")
_generator_nominal_power_v3(legacy_nominal)
nominal_columns = {
    row[1] for row in legacy_nominal.execute("PRAGMA table_info(generators)").fetchall()
}
assert "nominal_power_kw" in nominal_columns, nominal_columns
assert (
    legacy_nominal.execute("SELECT nominal_power_kw FROM generators WHERE id='g1'").fetchone()[0]
    is None
)
legacy_nominal.close()

# F00f: cadastral nominal kW persists, updates and can be explicitly cleared.
nominal_generator = db.create_generator(
    {
        "tag": "GEN-NOMINAL",
        "name": "Nominal test",
        "site": "Lab DB",
        "controller_type": "DSE",
        "controller_model": "DSE4520 MKII",
        "transport": "reverse_tcp",
        "host": "",
        "listen_port": 15049,
        "modbus_unit": 1,
        "nominal_power_kw": 450.0,
        "enabled": True,
    },
    actor="test",
)
assert nominal_generator["nominal_power_kw"] == 450.0, nominal_generator
nominal_generator = db.update_generator(
    nominal_generator["id"], {"nominal_power_kw": 500.0}, actor="test"
)
assert nominal_generator["nominal_power_kw"] == 500.0, nominal_generator
nominal_generator = db.update_generator(
    nominal_generator["id"], {"nominal_power_kw": None}, actor="test"
)
assert nominal_generator["nominal_power_kw"] is None, nominal_generator

# F01: the persistent users schema must accept the RBAC operator role.
with db.connect() as conn:
    users_sql = str(
        conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='users'"
        ).fetchone()[0]
    )
assert "'operador'" in users_sql, users_sql

# F00b: migration v2 rebuilds the legacy CHECK atomically and preserves child FKs.
legacy_path = root / "legacy-role.sqlite3"
legacy = sqlite3.connect(legacy_path)
legacy.row_factory = sqlite3.Row
legacy.execute("PRAGMA foreign_keys=ON")
legacy.executescript(
    """
    CREATE TABLE users(
        id TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        email TEXT NOT NULL UNIQUE COLLATE NOCASE,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL CHECK(role IN ('administrador','cadastro','visualizacao')),
        active INTEGER NOT NULL DEFAULT 1,
        last_access INTEGER,
        created_at INTEGER NOT NULL,
        updated_at INTEGER NOT NULL
    );
    CREATE TABLE sessions(
        token_hash TEXT PRIMARY KEY,
        user_id TEXT NOT NULL,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
    );
    INSERT INTO users(id,name,email,password_hash,role,active,last_access,created_at,updated_at)
    VALUES ('u1','Admin','admin-legacy@example.invalid','hash','administrador',1,NULL,1,1);
    INSERT INTO sessions(token_hash,user_id) VALUES ('s1','u1');
    """
)
_operator_role_v2(legacy)
legacy.execute(
    """INSERT INTO users(id,name,email,password_hash,role,active,last_access,created_at,updated_at)
       VALUES ('u2','Operador','operator-legacy@example.invalid','hash','operador',1,NULL,1,1)"""
)
assert legacy.execute("SELECT user_id FROM sessions WHERE token_hash='s1'").fetchone()[0] == "u1"
assert legacy.execute("PRAGMA foreign_key_check").fetchall() == []
legacy.close()

# F01: reset token must never escape through operational notification listings.
token = "super-secret-reset-token"
platform_store.enqueue_notification(
    "auth.password_reset",
    "email",
    destination="admin@example.invalid",
    subject="reset",
    body=f"https://example.invalid/reset-password?token={token}",
    payload={"token": token},
)
visible = platform_store.list_notifications(10)
assert visible and token not in repr(visible), visible
assert visible[0]["body"] == "Mensagem de segurança protegida"

# F02: setting up again cannot disable an already-enabled factor.
user = db.create_user(
    {
        "name": "Admin",
        "email": "admin@example.invalid",
        "password_hash": hash_password("StrongPass123!"),
        "role": "administrador",
        "active": True,
    },
    actor="test",
)
platform_store.set_totp(user["id"], protect_secret("JBSWY3DPEHPK3PXP"), True)
try:
    setup_totp(user)
except ValueError:
    pass
else:
    raise AssertionError("setup_totp replaced active factor")
assert platform_store.get_totp(user["id"])["enabled"] is True
try:
    disable_totp(user, totp_code("JBSWY3DPEHPK3PXP"), "SenhaIncorreta123")
except ValueError:
    pass
else:
    raise AssertionError("2FA disabled without current-password reauthentication")
assert platform_store.get_totp(user["id"])["enabled"] is True

# F03: administrative password replacement revokes all old sessions.
db.create_session("old-session-hash", user["id"], int(time.time()) + 3600)
assert db.get_session_user("old-session-hash") is not None
db.update_user(user["id"], {"password_hash": hash_password("AnotherPass123!")}, actor="test")
assert db.get_session_user("old-session-hash") is None

# F04: concurrent demotions cannot remove every active administrator.
second_admin = db.create_user(
    {
        "name": "Admin Secundário",
        "email": "admin2@example.invalid",
        "password_hash": hash_password("SecondAdmin123!"),
        "role": "administrador",
        "active": True,
    },
    actor="test",
)
admin_ids = [user["id"], second_admin["id"]]
admin_results = []
admin_lock = threading.Lock()


def demote_admin(user_id: str) -> None:
    try:
        db.update_user(user_id, {"role": "visualizacao"}, actor="race-test")
        outcome = "updated"
    except db.LastAdminError:
        outcome = "blocked"
    with admin_lock:
        admin_results.append(outcome)


admin_threads = [threading.Thread(target=demote_admin, args=(admin_id,)) for admin_id in admin_ids]
for thread in admin_threads:
    thread.start()
for thread in admin_threads:
    thread.join()
assert sorted(admin_results) == ["blocked", "updated"], admin_results
assert db.count_active_admins() == 1

remaining_admin = next(item for item in db.list_users() if item["role"] == "administrador" and item["active"])
try:
    db.delete_user(remaining_admin["id"], actor="race-test")
except db.LastAdminError:
    pass
else:
    raise AssertionError("último administrador ativo pôde ser excluído")

# F05: lifecycle industrial queue must allow only one active operation/executor.
life_generator = db.create_generator(
    {
        "tag": "GEN-LIFE",
        "name": "Lifecycle test",
        "site": "Lab DB",
        "controller_type": "COMAP",
        "controller_model": "InteliGen 200",
        "transport": "reverse_tcp",
        "host": "",
        "listen_port": 15050,
        "modbus_unit": 1,
        "enabled": True,
    },
    actor="test",
)

# Partial PATCHes must validate the final transport tuple, not just the changed field.
try:
    db.update_generator(life_generator["id"], {"listen_port": 60000}, actor="test")
except ValueError as exc:
    assert "RC_RAPID_LOCAL_OFFSET" in str(exc)
else:
    raise AssertionError("PATCH aceitou reverse TCP cujo listener local excede 65535")

try:
    db.update_generator(
        life_generator["id"],
        {"transport": "modbus_tcp_direct"},
        actor="test",
    )
except ValueError as exc:
    assert "host/IP" in str(exc)
else:
    raise AssertionError("PATCH aceitou transporte TCP direto sem host")

assert db.get_generator(life_generator["id"])["transport"] == "reverse_tcp"
assert db.get_generator(life_generator["id"])["listen_port"] == 15050

enqueue_results = []
enqueue_lock = threading.Lock()


def enqueue_lifecycle(operation_id: str) -> None:
    try:
        item = platform_store.enqueue_lifecycle_operation(
            operation_id,
            life_generator["id"],
            "provision",
            {"confirmation": "PROVISIONAR"},
            "test",
            "administrador",
        )
        result = ("queued", item["operationId"])
    except ValueError:
        result = ("blocked", operation_id)
    with enqueue_lock:
        enqueue_results.append(result)


enqueue_threads = [
    threading.Thread(target=enqueue_lifecycle, args=("life-op-a",)),
    threading.Thread(target=enqueue_lifecycle, args=("life-op-b",)),
]
for thread in enqueue_threads:
    thread.start()
for thread in enqueue_threads:
    thread.join()
assert sorted(item[0] for item in enqueue_results) == ["blocked", "queued"], enqueue_results

claim_results = []
claim_lock = threading.Lock()


def claim_lifecycle() -> None:
    item = platform_store.claim_lifecycle_operation()
    with claim_lock:
        claim_results.append(item)


claim_threads = [threading.Thread(target=claim_lifecycle) for _ in range(2)]
for thread in claim_threads:
    thread.start()
for thread in claim_threads:
    thread.join()
claimed = [item for item in claim_results if item]
assert len(claimed) == 1, claim_results
claimed_id = claimed[0]["operationId"]
with db.connect() as conn:
    conn.execute(
        "UPDATE lifecycle_operations SET status='failed',error='synthetic timeout' WHERE id=?",
        (claimed_id,),
    )
assert platform_store.finish_lifecycle_operation(claimed_id, result={"late": True}) is False
assert platform_store.get_lifecycle_operation(claimed_id)["status"] == "failed"

# F06: scheduled jobs are leased atomically; two workers cannot execute one due job.
scheduled = platform_store.upsert_scheduler_job(
    {
        "id": "job-race",
        "name": "Race-safe scheduler",
        "kind": "notification",
        "interval_seconds": 60,
        "payload": {},
        "enabled": True,
        "next_run": 1,
    },
    "test",
)
scheduler_claims = []
scheduler_lock = threading.Lock()


def claim_scheduler() -> None:
    rows = platform_store.claim_scheduler_jobs({"notification"}, limit=1)
    with scheduler_lock:
        scheduler_claims.append(rows)


scheduler_threads = [threading.Thread(target=claim_scheduler) for _ in range(2)]
for thread in scheduler_threads:
    thread.start()
for thread in scheduler_threads:
    thread.join()
claimed_jobs = [row for batch in scheduler_claims for row in batch]
assert len(claimed_jobs) == 1, scheduler_claims
first_scheduler_claim = claimed_jobs[0]
assert first_scheduler_claim["id"] == scheduled["id"]
assert first_scheduler_claim["claim_token"]
public_scheduler = next(item for item in platform_store.list_scheduler_jobs() if item["id"] == scheduled["id"])
assert "claim_token" not in public_scheduler

# Simulate an expired lease being reclaimed while the original worker is late.
with db.connect() as conn:
    conn.execute("UPDATE scheduler_jobs SET next_run=1 WHERE id=?", (scheduled["id"],))
second_scheduler_claims = platform_store.claim_scheduler_jobs({"notification"}, limit=1)
assert len(second_scheduler_claims) == 1
second_scheduler_claim = second_scheduler_claims[0]
assert second_scheduler_claim["claim_token"] != first_scheduler_claim["claim_token"]

assert (
    platform_store.complete_scheduler_job(
        scheduled["id"],
        "late stale worker",
        first_scheduler_claim["claim_token"],
    )
    is False
)
assert (
    platform_store.complete_scheduler_job(
        scheduled["id"],
        "OK synthetic",
        second_scheduler_claim["claim_token"],
    )
    is True
)

# F07: automation uses effective stale status and retries a failed edge.
stale = {"id": "g", "tag": "GEN-AUTO", "status": "alerta", "telemetryStale": True}
assert automation_engine._condition({"type": "generator_offline", "value": "GEN-AUTO"}, [stale])[0]
assert not automation_engine._condition({"type": "generator_online", "value": "GEN-AUTO"}, [stale])[0]
assert not automation_engine._condition({"type": "generator_alert", "value": "GEN-AUTO"}, [stale])[0]

rule = ops_store.create_rule(
    {
        "name": "Retry edge",
        "trigger": "generator_online:GEN-AUTO",
        "action": "notify:panel",
    },
    "test",
)
automation_engine.approve_rule(rule["id"], "test")
automation_engine.set_rule_enabled(rule["id"], True, "test")
original_overlay = automation_engine.overlay_generators
original_execute = automation_engine._execute
attempts = {"count": 0}


def synthetic_overlay(_generators):
    return [{"id": "g", "tag": "GEN-AUTO", "status": "online", "telemetryStale": False}]


def flaky_execute(_action, _rule, _generator):
    attempts["count"] += 1
    if attempts["count"] == 1:
        raise RuntimeError("synthetic delivery failure")
    return "notify:panel"


automation_engine.overlay_generators = synthetic_overlay
automation_engine._execute = flaky_execute
try:
    assert automation_engine.process_rules() == 0
    with db.connect() as conn:
        state = conn.execute(
            "SELECT last_value FROM automation_state WHERE rule_id=?", (rule["id"],)
        ).fetchone()
    assert state is None or state["last_value"] != "1"
    assert automation_engine.process_rules() == 1
    assert automation_engine.process_rules() == 0
    assert attempts["count"] == 2
finally:
    automation_engine.overlay_generators = original_overlay
    automation_engine._execute = original_execute

# F07: external API read allowlist is least-privilege when configured.
api_generator = {"id": "gen-api-1", "tag": "GEN-API-1"}
assert _token_allows_generator({"allowed_generators": []}, api_generator)
assert _token_allows_generator({"allowed_generators": ["gen-api-1"]}, api_generator)
assert _token_allows_generator({"allowed_generators": ["GEN-API-1"]}, api_generator)
assert not _token_allows_generator({"allowed_generators": ["GEN-OTHER"]}, api_generator)

# F08: report artifacts may never escape DATA_DIR/reports.
reports_dir = data / "reports"
reports_dir.mkdir(parents=True, exist_ok=True)
inside_report = reports_dir / "safe.csv"
inside_report.write_text("ok", encoding="utf-8")
assert safe_report_artifact_path(inside_report) == inside_report.resolve()
outside_report = root / "outside.csv"
outside_report.write_text("secret", encoding="utf-8")
try:
    safe_report_artifact_path(outside_report)
except ValueError:
    pass
else:
    raise AssertionError("artefato de relatório fora de DATA_DIR/reports foi aceito")

delete_probe = reports_dir / "delete-probe.csv"
delete_probe.write_text("keep-until-commit", encoding="utf-8")
staged_probe = _stage_file_delete(delete_probe, reports_dir)
assert staged_probe is not None
assert not delete_probe.exists() and staged_probe[1].exists()
_restore_staged_file(staged_probe)
assert delete_probe.read_text(encoding="utf-8") == "keep-until-commit"

staged_probe = _stage_file_delete(delete_probe, reports_dir)
assert staged_probe is not None
_finish_staged_file(staged_probe)
assert not delete_probe.exists() and not staged_probe[1].exists()

try:
    _stage_file_delete(outside_report, reports_dir)
except ValueError:
    pass
else:
    raise AssertionError("deleção protegida aceitou arquivo fora do diretório")

# F09: operational records must reject dangling generator references before SQLite.
for create_invalid_reference in (
    lambda: ops_store.create_work_order(
        {
            "generator_id": "missing-generator",
            "gen": "",
            "site": "",
            "type": "Preventiva",
            "due": 0,
            "tech": "",
            "status": "Planejada",
            "description": "",
        },
        "test",
    ),
    lambda: ops_store.create_agenda(
        {
            "title": "Inspeção",
            "when": "amanhã",
            "site": "",
            "generator_id": "missing-generator",
            "kind": "manual",
        },
        "test",
    ),
):
    try:
        create_invalid_reference()
    except ValueError:
        pass
    else:
        raise AssertionError("referência a gerador inexistente foi aceita")

# F10: field-device inventory must reject dangling site/generator references.
for payload in (
    {"kind": "modem", "name": "Bad generator", "generator_id": "missing-generator"},
    {"kind": "gateway", "name": "Bad site", "site_id": "missing-site"},
):
    try:
        platform_store.create_field_device(payload, "test")
    except ValueError:
        pass
    else:
        raise AssertionError("field-device com referência inexistente foi aceito")

# F11: authentication throttles must not lose concurrent increments.
race_key = platform_store.login_key("race@example.invalid", "192.0.2.44")
login_threads = [
    threading.Thread(
        target=platform_store.record_login_failure,
        args=(race_key,),
        kwargs={"max_failures": 5, "lock_seconds": 900},
    )
    for _ in range(5)
]
for thread in login_threads:
    thread.start()
for thread in login_threads:
    thread.join()
allowed, _retry = platform_store.login_allowed(race_key, max_failures=5, lock_seconds=900)
assert allowed is False
with db.connect() as conn:
    login_row = conn.execute(
        "SELECT failures,locked_until FROM login_attempts WHERE attempt_key=?",
        (race_key,),
    ).fetchone()
assert int(login_row["failures"]) == 5
assert int(login_row["locked_until"]) > int(time.time())

reset_email = "reset-race@example.invalid"
reset_ip = "192.0.2.45"
reset_results = []
reset_lock = threading.Lock()


def reset_race() -> None:
    result = platform_store.password_reset_allowed(
        reset_email,
        reset_ip,
        max_per_window=10,
        account_max_per_window=10,
        cooldown_seconds=0,
    )
    with reset_lock:
        reset_results.append(result)


reset_threads = [threading.Thread(target=reset_race) for _ in range(4)]
for thread in reset_threads:
    thread.start()
for thread in reset_threads:
    thread.join()
assert reset_results == [True, True, True, True]
pair_key = "pair:" + hashlib.sha256(f"{reset_email}|{reset_ip}".encode()).hexdigest()
account_key = "acct:" + hashlib.sha256(reset_email.encode()).hexdigest()
with db.connect() as conn:
    pair_count = conn.execute(
        "SELECT requests FROM password_reset_requests WHERE request_key=?", (pair_key,)
    ).fetchone()[0]
    account_count = conn.execute(
        "SELECT requests FROM password_reset_requests WHERE request_key=?", (account_key,)
    ).fetchone()[0]
assert int(pair_count) == 4
assert int(account_count) == 4

# F10: reset flow is throttled independently from login.
assert platform_store.password_reset_allowed("target@example.invalid", "192.0.2.10") is True
assert platform_store.password_reset_allowed("target@example.invalid", "192.0.2.10") is False

# F05/F11: spreadsheet text is neutralized, numeric telemetry stays numeric, report is a snapshot.
report = {"id": "rep-audit", "name": "=1+1", "period": "Ontem", "format": "XLSX"}
generator = {
    "tag": "GEN001",
    "site": "=2+2",
    "status": "online",
    "controller": "ComAp",
    "availableMetrics": ["rpm", "fuel_level"],
    "definedMetrics": ["rpm", "fuel_level"],
    "telemetryStale": False,
    "metricUnits": {"fuel_level": "L"},
    "rpm": 1500,
    "fuelLevel": 1091,
}
artifact = generate_report(report, [generator])
wb = load_workbook(artifact["path"], data_only=False)
ws = wb.active
assert ws["A1"].data_type != "f" and str(ws["A1"].value).startswith("'")
assert ws["B6"].data_type != "f" and str(ws["B6"].value).startswith("'")
assert ws["E6"].value == 1500 and ws["E6"].data_type == "n"
assert ws["I6"].value == 1091 and ws["J6"].value == "L"
assert ws["A2"].value == "Tipo" and ws["B2"].value == "Fotografia operacional"

stale_report = {"id": "rep-stale", "name": "Stale", "period": "Agora", "format": "XLSX"}
stale_generator = {
    **generator,
    "definedMetrics": [],
    "telemetryStale": True,
    "rpm": 1800,
    "fuelLevel": 1200,
}
stale_artifact = generate_report(stale_report, [stale_generator])
stale_wb = load_workbook(stale_artifact["path"], data_only=False)
stale_ws = stale_wb.active
assert stale_ws["E6"].value is None
assert stale_ws["I6"].value is None
assert stale_ws["J6"].value is None

tracked_report = ops_store.create_report(
    {"name": "Atomic", "period": "Agora", "format": "CSV"},
    "test",
)
assert tracked_report["status"] == "Gerando"
tracked_artifact = generate_report(tracked_report, [generator])
tracked_after = next(
    item for item in ops_store.list_reports() if item["id"] == tracked_report["id"]
)
assert tracked_after["status"] == "Pronto"
assert tracked_artifact["path"].is_file()
assert not list((data / "reports").glob(f".{tracked_report['id']}.*.tmp"))

bad_report = ops_store.create_report(
    {"name": "Bad", "period": "Agora", "format": "INVALID"},
    "test",
)
try:
    generate_report(bad_report, [generator])
except ValueError:
    pass
else:
    raise AssertionError("gerador de relatório aceitou formato inválido")
bad_after = next(item for item in ops_store.list_reports() if item["id"] == bad_report["id"])
assert bad_after["status"] == "Falha"
assert platform_store.get_report_artifact(bad_report["id"]) is None

# F06: clearing and reopening the same alarm starts a fresh escalation occurrence.
policy = industrial_store.create_escalation_policy(
    {
        "name": "comm",
        "severity": "fault",
        "after_seconds": 0,
        "channel": "panel",
        "repeat_seconds": 0,
        "max_repeats": 1,
    },
    "test",
)
offline = [{"id": "g1", "tag": "GEN001", "status": "offline", "lastError": "loss"}]
online = [{"id": "g1", "tag": "GEN001", "status": "online"}]
assert industrial_store.process_escalations(offline) == 1
industrial_store.process_escalations(online)
assert industrial_store.process_escalations(offline) == 1
with db.connect() as conn:
    count = conn.execute(
        "SELECT COUNT(*) FROM notification_queue WHERE event_type='industrial.alarm.escalation'"
    ).fetchone()[0]
assert count == 2, count

# A mesma ocorrência de alarme não pode ser escalada duas vezes por workers concorrentes.
race_policy = industrial_store.create_escalation_policy(
    {
        "name": "comm-race",
        "severity": "fault",
        "after_seconds": 0,
        "channel": "panel",
        "repeat_seconds": 0,
        "max_repeats": 1,
    },
    "test",
)
race_offline = [
    {"id": "g-race", "tag": "GEN-RACE", "status": "offline", "lastError": "race-loss"}
]
race_results = []
race_errors = []
race_lock = threading.Lock()


def escalate_race() -> None:
    try:
        result = industrial_store.process_escalations(race_offline)
    except Exception as exc:
        with race_lock:
            race_errors.append(repr(exc))
        return
    with race_lock:
        race_results.append(result)


race_threads = [threading.Thread(target=escalate_race) for _ in range(2)]
for thread in race_threads:
    thread.start()
for thread in race_threads:
    thread.join()
assert not race_errors, race_errors
with db.connect() as conn:
    race_rows = conn.execute(
        "SELECT payload_json FROM notification_queue "
        "WHERE event_type='industrial.alarm.escalation'"
    ).fetchall()
race_matches = [
    row
    for row in race_rows
    if json.loads(row["payload_json"] or "{}").get("policyId") == race_policy["id"]
    and json.loads(row["payload_json"] or "{}").get("generatorId") == "g-race"
]
assert len(race_matches) == 1, (race_results, race_matches)

# F07: an abandoned notification lease is reclaimed, and a late worker is fenced out.
queue_id = platform_store.enqueue_notification("test.lease", "panel", body="lease")
claimed = platform_store.claim_due_notifications(20, lease_seconds=30)
first_claim = next(item for item in claimed if item["id"] == queue_id)
assert first_claim["claim_token"]
assert "claim_token" not in platform_store.list_notifications(50)[0]

with db.connect() as conn:
    conn.execute(
        "UPDATE notification_queue SET status='sending',updated_at=? WHERE id=?",
        (int(time.time()) - 120, queue_id),
    )
reclaimed = platform_store.claim_due_notifications(20, lease_seconds=30)
second_claim = next(item for item in reclaimed if item["id"] == queue_id)
assert second_claim["claim_token"]
assert second_claim["claim_token"] != first_claim["claim_token"]

assert (
    platform_store.finish_notification(
        queue_id,
        "panel",
        "",
        True,
        "late stale worker",
        first_claim["claim_token"],
    )
    is False
)
with db.connect() as conn:
    still_sending = conn.execute(
        "SELECT status,claim_token,attempts FROM notification_queue WHERE id=?",
        (queue_id,),
    ).fetchone()
assert still_sending["status"] == "sending"
assert still_sending["claim_token"] == second_claim["claim_token"]
assert int(still_sending["attempts"]) == 0

assert (
    platform_store.finish_notification(
        queue_id,
        "panel",
        "",
        True,
        "active worker",
        second_claim["claim_token"],
    )
    is True
)
with db.connect() as conn:
    finished = conn.execute(
        "SELECT status,claim_token,attempts FROM notification_queue WHERE id=?",
        (queue_id,),
    ).fetchone()
assert finished["status"] == "sent"
assert finished["claim_token"] == ""
assert int(finished["attempts"]) == 1

# F08: stale telemetry must be fail-closed in backend summaries and alarms.
summary = dashboard(
    [
        {"status": "online", "telemetryStale": False},
        {"status": "alerta", "telemetryStale": True},
        {"status": "nao_configurado", "telemetryStale": True},
    ]
)
assert summary["online"] == 1, summary
assert summary["alerts"] == 0, summary
assert summary["offline"] == 1, summary
assert summary["notConfigured"] == 1, summary

industrial_store.refresh_observed_alarms(
    [
        {
            "id": "g-stale",
            "tag": "GEN-STALE",
            "status": "alerta",
            "telemetryStale": True,
            "lastError": "reader timeout",
            "definedMetrics": [],
        }
    ]
)
with db.connect() as conn:
    stale_alarm = conn.execute(
        "SELECT code,source,active FROM industrial_alarms WHERE alarm_key=?",
        ("comm:g-stale",),
    ).fetchone()
assert stale_alarm is not None
assert stale_alarm["code"] == "COMM_LOSS"
assert stale_alarm["source"] == "derived.communication"
assert int(stale_alarm["active"]) == 1

# F09: downsampling preserves the complete window and a strict bound.
for size in (2001, 2880, 3999, 10000):
    points = [{"timestamp": str(i), "value": i, "stat": 1} for i in range(size)]
    sampled = _downsample_points(points, 2000)
    assert len(sampled) == 2000
    assert sampled[0]["value"] == 0
    assert sampled[-1]["value"] == size - 1

# F10: lifecycle operations are durable, exclusive per generator and idempotent by ID.
operation_payload = {"confirmation": "PROVISION", "operationId": "op-audit-0001"}
queued = platform_store.enqueue_lifecycle_operation(
    "op-audit-0001",
    "g-operation",
    "provision",
    operation_payload,
    "admin@example.invalid",
    "administrador",
)
assert queued["status"] == "queued"
claimed_operation = platform_store.claim_lifecycle_operation()
assert claimed_operation and claimed_operation["operationId"] == "op-audit-0001"
assert claimed_operation["status"] == "running"
try:
    platform_store.enqueue_lifecycle_operation(
        "op-audit-0002",
        "g-operation",
        "deprovision",
        {"confirmation": "DEPROVISION", "operationId": "op-audit-0002"},
        "admin@example.invalid",
        "administrador",
    )
except ValueError:
    pass
else:
    raise AssertionError("concurrent lifecycle operation accepted for same generator")
platform_store.finish_lifecycle_operation("op-audit-0001", {"ok": True})
finished = platform_store.get_lifecycle_operation("op-audit-0001", "g-operation")
assert finished and finished["status"] == "succeeded"
assert finished["result"]["ok"] is True
replayed = platform_store.enqueue_lifecycle_operation(
    "op-audit-0001",
    "g-operation",
    "provision",
    operation_payload,
    "admin@example.invalid",
    "administrador",
)
assert replayed["status"] == "succeeded"

# Observability: each supervised worker is individually visible and queue health is bounded.
for worker_name in platform_store.EXPECTED_WORKERS:
    platform_store.touch_worker_heartbeat(worker_name, "ok", "regression")
workers = platform_store.worker_health(stale_after=45)
assert len(workers) == 4 and all(item["healthy"] for item in workers), workers
queue_health = platform_store.queue_health()
assert queue_health["staleNotificationClaims"] == 0, queue_health
assert queue_health["staleLifecycleOperations"] == 0, queue_health
assert queue_health["healthy"] is True, queue_health


# Reverse TCP peer history must aggregate metadata without storing transport payloads.
first_peer = traffic_store.record_bridge_peer(
    15001,
    "203.0.113.10",
    accepted=True,
    reason="conexão aceita",
    now=1_700_000_000,
)
assert first_peer["acceptedCount"] == 1 and first_peer["rejectedCount"] == 0, first_peer
second_peer = traffic_store.record_bridge_peer(
    15001,
    "203.0.113.10",
    accepted=False,
    reason="origem fora da allowlist",
    now=1_700_000_060,
)
assert second_peer["acceptedCount"] == 1 and second_peer["rejectedCount"] == 1, second_peer
assert second_peer["firstSeenAt"] == 1_700_000_000
assert second_peer["lastSeenAt"] == 1_700_000_060
assert second_peer["lastDecision"] == "rejected"
peer_rows = traffic_store.list_bridge_peers(10, remote_port=15001)
assert peer_rows and peer_rows[0]["remoteIp"] == "203.0.113.10", peer_rows
assert "payload" not in peer_rows[0] and "frame" not in peer_rows[0], peer_rows[0]
traffic_store.record_bridge_peer(
    15002,
    "203.0.113.11",
    accepted=True,
    reason="conexão aceita",
    now=1_700_000_060 + 91 * 86400,
)
retained = traffic_store.list_bridge_peers(20)
assert all(item["remoteIp"] != "203.0.113.10" for item in retained), retained
assert any(item["remoteIp"] == "203.0.113.11" for item in retained), retained

# Reverse TCP security posture must remain visible to diagnostics.
bridge_status = root / "bridge-status.json"
bridge_status.write_text(
    '{"updatedAt": %d, "security": {"peerAllowlistEnabled": false, "peerAllowlistRequired": false}, "ports": []}'
    % int(time.time())
)
previous_bridge_status = diagnostics.BRIDGE_STATUS_FILE
try:
    diagnostics.BRIDGE_STATUS_FILE = bridge_status
    status = diagnostics._bridge_runtime_status()
finally:
    diagnostics.BRIDGE_STATUS_FILE = previous_bridge_status
assert status["security"]["peerAllowlistEnabled"] is False, status

print("Audit regressions: OK")


# F12: production Controller Packs use schema v4 and every enabled command has a contract.
from app.controller_library import list_controller_packs  # noqa: E402

for pack in list_controller_packs():
    if pack.get("lifecycle") != "production":
        continue
    assert pack.get("schema") == 4, pack.get("packId")
    capabilities = pack.get("capabilities") or {}
    contracts = pack.get("commands") or {}
    for action in (
        "start",
        "stop",
        "auto",
        "manual",
        "test",
        "mcb_open",
        "mcb_close",
        "gcb_open",
        "gcb_close",
        "paralleling",
    ):
        if capabilities.get(action):
            assert action in contracts, (pack.get("packId"), action)

# F13: production refuses experimental control flags.
from app.production_guard import validate_production_runtime  # noqa: E402

previous_environment = os.environ.get("RC_ENVIRONMENT")
previous_dse_lab = os.environ.get("RC_ENABLE_DSE_LAB_CONTROL")
try:
    os.environ["RC_ENVIRONMENT"] = "production"
    os.environ["RC_ENABLE_DSE_LAB_CONTROL"] = "1"
    try:
        validate_production_runtime()
    except RuntimeError:
        pass
    else:
        raise AssertionError("production accepted RC_ENABLE_DSE_LAB_CONTROL=1")
finally:
    if previous_environment is None:
        os.environ.pop("RC_ENVIRONMENT", None)
    else:
        os.environ["RC_ENVIRONMENT"] = previous_environment
    if previous_dse_lab is None:
        os.environ.pop("RC_ENABLE_DSE_LAB_CONTROL", None)
    else:
        os.environ["RC_ENABLE_DSE_LAB_CONTROL"] = previous_dse_lab

# F14: the composed API must not expose duplicate method/path contracts.
import re  # noqa: E402
from collections import defaultdict  # noqa: E402
from fastapi.routing import APIRoute  # noqa: E402
from app.main import app  # noqa: E402


def _walk_routes(routes):
    for route in routes:
        if isinstance(route, APIRoute):
            yield route
        elif type(route).__name__ == "_IncludedRouter":
            yield from _walk_routes(route.original_router.routes)


contracts = defaultdict(list)
for route in _walk_routes(app.routes):
    normalized = re.sub(r"\{[^}]+\}", "{}", route.path)
    for method in route.methods or []:
        contracts[(method, normalized)].append(route.name)
duplicates = {key: names for key, names in contracts.items() if len(names) > 1}
assert not duplicates, duplicates

print("Production hardening regressions: OK")
